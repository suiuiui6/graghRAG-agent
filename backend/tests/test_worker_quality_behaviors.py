"""Worker behavior checks with real temporary files and explicit provider doubles.

These are local component tests, not live MinerU/Ollama integration tests.
"""
from __future__ import annotations

import json
import sys
from pathlib import Path
from types import SimpleNamespace

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import worker
from doc_processor import Section
from errors import (
    ErrorCode, PipelineStage, StageBridgeError, StageExtractionError,
    StageGroundingError, StageKGBuildError, StageMinerUError, classify_external_error,
)


@pytest.fixture(autouse=True)
def isolated_worker(monkeypatch, tmp_path):
    """Every worker-generated path is under pytest's unique sandbox."""
    monkeypatch.setenv("PYTHON_DOTENV_DISABLED", "1")
    monkeypatch.setenv("HF_HUB_OFFLINE", "1")
    monkeypatch.setenv("TRANSFORMERS_OFFLINE", "1")
    monkeypatch.setattr(worker, "__file__", str(tmp_path / "worker.py"))
    monkeypatch.setattr(worker, "UPLOADS_DIR", str(tmp_path / "uploads"))
    import requests

    def deny_network(*args, **kwargs):
        raise AssertionError("Unstubbed provider request is forbidden")

    monkeypatch.setattr(requests.sessions.Session, "request", deny_network)


def _chunks(count):
    return [SimpleNamespace(text=f"设备与人员关系 {index}") for index in range(count)]


def _provider(monkeypatch, responses):
    """Only the HTTP boundary is doubled; extraction and JSON writes stay real."""
    import requests
    queue = iter(responses)
    calls = []

    def post(url, **kwargs):
        calls.append((url, kwargs))
        payload = next(queue)
        if isinstance(payload, Exception):
            raise payload
        return SimpleNamespace(
            raise_for_status=lambda: None,
            json=lambda: {"message": {"content": json.dumps(payload, ensure_ascii=False)}},
        )

    monkeypatch.setattr(requests, "post", post)
    return calls


def _graph_payload():
    return {
        "entities": [
            {"name": "班组", "type": "机构"}, {"name": "设备"},
            {"name": "班组", "type": "其他"}, {"name": ""},
        ],
        "relations": [
            {"subject": "班组", "object": "设备", "predicate": "监管"},
            {"subject": "班组", "object": "设备", "predicate": "监管"},
            {"subject": "班组", "object": "未知"},
            {"subject": "", "object": "设备"},
        ],
    }


def test_extraction_deduplicates_nodes_edges_and_persists_task_owned_graph(monkeypatch, tmp_path):
    calls = _provider(monkeypatch, [_graph_payload(), _graph_payload()])
    outcome = worker._extract_and_build_kg(_chunks(6), "owned-task", "设备.docx")
    assert (outcome.node_count, outcome.edge_count) == (2, 1)
    assert (outcome.batches_total, outcome.batches_failed, outcome.partial) == (2, 0, False)
    graph = tmp_path / "output/kg/doc_owned-task"
    nodes = json.loads((graph / "nodes.json").read_text(encoding="utf-8"))["nodes"]
    edges = json.loads((graph / "edges.json").read_text(encoding="utf-8"))["edges"]
    assert [(node["label"], node["entity_type"]) for node in nodes] == [("班组", "机构"), ("设备", "其他")]
    assert edges == [{"source_node_id": nodes[0]["node_id"], "target_node_id": nodes[1]["node_id"], "relation_type": "监管"}]
    assert len(calls) == 2
    assert calls[0][1]["json"]["stream"] is False
    assert calls[0][1]["timeout"] == 120


def test_failed_batch_recovers_to_partial_graph_with_first_root_cause(monkeypatch, tmp_path):
    import requests
    _provider(monkeypatch, [requests.Timeout("provider deadline"), _graph_payload()])
    outcome = worker._extract_and_build_kg(_chunks(6), "recover-task", "设备.docx")
    assert outcome.partial and not outcome.all_failed
    assert (outcome.batches_total, outcome.batches_failed) == (2, 1)
    assert outcome.first_error.code == ErrorCode.PIPELINE_EXTRACTION_TIMEOUT
    assert outcome.first_error.context.task_id == "recover-task"
    assert json.loads((tmp_path / "output/kg/doc_recover-task/nodes.json").read_text(encoding="utf-8"))["nodes"][0]["label"] == "班组"


def test_all_failed_batches_raise_with_counts_and_do_not_create_graph(monkeypatch, tmp_path):
    import requests
    _provider(monkeypatch, [requests.Timeout("first failure"), ValueError("invalid JSON")])
    with pytest.raises(StageExtractionError) as captured:
        worker._extract_and_build_kg(_chunks(6), "failed-task", "设备.docx")
    error = captured.value
    assert error.code == ErrorCode.PIPELINE_EXTRACTION_EMPTY
    assert error.context.details == {"batches_total": 2, "batches_failed": 2, "root_error": "PIPELINE_EXTRACTION_TIMEOUT"}
    assert error.context.retryable is True
    assert error.__cause__.code == ErrorCode.PIPELINE_EXTRACTION_TIMEOUT
    assert not (tmp_path / "output").exists()


@pytest.mark.parametrize("chunks,expected_batches", [([], 0), ([SimpleNamespace(text="  "), object()], 0), (_chunks(1), 1)])
def test_empty_input_or_empty_entities_does_not_claim_failed_extraction(monkeypatch, tmp_path, chunks, expected_batches):
    calls = _provider(monkeypatch, [{"entities": [], "relations": []}])
    result = worker._extract_and_build_kg(chunks, "empty-task", "empty.docx")
    assert (result.node_count, result.edge_count, result.batches_total) == (0, 0, expected_batches)
    assert result.first_error is None and not result.partial and not result.all_failed
    assert len(calls) == expected_batches
    assert not (tmp_path / "output").exists()


def test_extraction_limits_batch_budget_and_prompt_length(monkeypatch):
    calls = _provider(monkeypatch, [{"entities": []}] * 10)
    chunks = [SimpleNamespace(text="甲" * 5000)] + _chunks(50)
    outcome = worker._extract_and_build_kg(chunks, "bounded-task", "large.docx")
    assert outcome.batches_total == 10 and len(calls) == 10
    first_prompt = calls[0][1]["json"]["messages"][0]["content"]
    assert first_prompt.split("文本:\n", 1)[1] == "甲" * 4000
    assert "设备与人员关系 49" not in calls[-1][1]["json"]["messages"][0]["content"]


@pytest.mark.parametrize("task_id", ["", "../other", "a/b", "a\\b", ".hidden", "NUL", "con", "COM1", "LPT9", "a" * 129, "任务", "x:", "x."])
def test_mineru_rejects_unsafe_task_paths_before_provider_submission(monkeypatch, tmp_path, task_id):
    def no_submission(*args, **kwargs):
        pytest.fail("Unsafe task identifier reached the upload provider")
    monkeypatch.setattr(worker, "submit_url_task", no_submission)
    monkeypatch.setattr(worker, "submit_local_file", no_submission)
    for call in (
        lambda: worker._call_mineru_api("https://example.invalid/a.pdf", "test-only", task_id),
        lambda: worker._call_mineru_local_file(str(tmp_path / "a.pdf"), "test-only", task_id),
    ):
        with pytest.raises(StageMinerUError) as captured:
            call()
        assert captured.value.code == ErrorCode.PIPELINE_MINERU_PARSE
        assert captured.value.context.retryable is False
    assert not (tmp_path / "output").exists()


@pytest.mark.parametrize("local", [True, False])
def test_mineru_transports_keep_complete_identifier_and_propagate_options(monkeypatch, tmp_path, local):
    source = tmp_path / "original.pdf"
    source.write_bytes(b"local PDF boundary fixture")
    task_id = "ingest_sameprefix_1234567890"
    options = {"language": "ch", "enable_table": False}
    calls = []

    def submit(*args, **kwargs):
        calls.append(("submit", args, kwargs))
        return "handle"

    def poll(handle, token, **kwargs):
        assert (handle, token, kwargs) == ("handle", "test-only", {"task_id": task_id})
        calls.append(("poll",))
        return "https://example.invalid/parse.zip"

    def download(url, output_dir, **kwargs):
        assert url == "https://example.invalid/parse.zip"
        assert kwargs == {"task_id": task_id}
        target = Path(output_dir)
        target.mkdir(parents=True)
        (target / "full.md").write_text("本次任务", encoding="utf-8")
        calls.append(("download",))
        return str(target)

    monkeypatch.setattr(worker, "submit_local_file" if local else "submit_url_task", submit)
    monkeypatch.setattr(worker, "poll_until_done", poll)
    monkeypatch.setattr(worker, "download_and_extract", download)
    function = worker._call_mineru_local_file if local else worker._call_mineru_api
    output = function(str(source) if local else "https://example.invalid/a.pdf", "test-only", task_id, options)
    assert Path(output) == tmp_path / "output/mineru" / task_id
    assert Path(output, "full.md").read_text(encoding="utf-8") == "本次任务"
    assert [call[0] for call in calls] == ["submit", "poll", "download"]
    if local:
        assert calls[0][2] == {"local_path": str(source), "filename": "original.pdf", "task_id": task_id, "ingest_options": options}
    else:
        assert calls[0][1] == ("https://example.invalid/a.pdf", "test-only", task_id, options)


def test_sidecar_csv_batches_rows_and_preserves_offsets_and_cell_width(tmp_path):
    path = tmp_path / "evidence.csv"
    path.write_text("\ufeff场景,,备注\nA,1\nB,2,有效,discard\n" + "C,3,记录\n" * 100, encoding="utf-8")
    sections, end = worker._sidecar_table_sections(path, base_title="", char_offset=17)
    assert len(sections) == 2
    assert sections[0].table_json["headers"] == ["场景", "col_2", "备注"]
    assert sections[0].table_json["rows"][:2] == [["A", "1", ""], ["B", "2", "有效"]]
    assert [len(section.table_json["rows"]) for section in sections] == [100, 2]
    assert sections[0].table_json["caption"] == "evidence.csv (rows 1-100)"
    assert sections[1].table_json["caption"] == "evidence.csv (rows 101-102)"
    assert sections[0].char_start == 17
    assert sections[1].char_start == sections[0].char_end
    assert end == sections[-1].char_end
    assert all(section.title == "evidence" for section in sections)


@pytest.mark.parametrize("body", ["", "head\n", "head\n,\n", "head\n,ignored\n"])
def test_sidecar_csv_without_valid_rows_keeps_offset(tmp_path, body):
    path = tmp_path / "empty.csv"
    path.write_text(body, encoding="utf-8")
    sections, offset = worker._sidecar_table_sections(path, base_title="空", char_offset=9)
    assert sections == [] and offset == 9


def test_sidecar_reference_discovery_prefers_mineru_and_deduplicates_source_directory(tmp_path):
    mineru = tmp_path / "mineru"
    source_dir = tmp_path / "source"
    mineru.mkdir()
    source_dir.mkdir()
    (mineru / "evidence.csv").write_text("来源\n解析输出\n", encoding="utf-8")
    (source_dir / "evidence.csv").write_text("来源\n原文附件\n", encoding="utf-8")
    (source_dir / "note.md").write_text("\ufeff 附件正文 ", encoding="utf-8")
    (mineru / "empty.txt").write_text("  ", encoding="utf-8")
    references = [Section(title="附件", level=2, text="evidence.csv evidence.csv note.md empty.txt missing.json", content_type="text")]
    sections, end = worker._load_sidecar_sections(str(mineru), source_path=str(source_dir / "source.pdf"), reference_sections=references, base_title="附件", char_offset=5)
    assert len(sections) == 2
    assert sections[0].table_json["rows"] == [["解析输出"]]
    assert sections[1].text == "note.md\n附件正文"
    assert sections[1].char_start == sections[0].char_end
    assert end == sections[1].char_end
    assert worker._discover_support_artifact_paths(str(mineru), source_path=str(mineru / "source.pdf"), references=["evidence.csv", "evidence.csv"]) == [mineru / "evidence.csv"]


@pytest.mark.parametrize("item,expected", [
    ({"table_body": ""}, None),
    ({"table_body": "<tr><td>Only header</td></tr>"}, None),
    ({"table_body": "<tr><td>H</td><td></td></tr><tr><td>A &amp; B</td></tr><tr><td>C</td><td> D </td><td>E</td></tr>", "table_caption": ["表", "一"]},
     {"headers": ["H", "col_2", "col_3"], "rows": [["A & B", "", ""], ["C", "D", "E"]], "caption": "表 一"}),
])
def test_mineru_html_tables_keep_normalized_schema(item, expected):
    assert worker._mineru_table_item_to_table_json(item) == expected


def _pdf_fixture(tmp_path):
    directory = tmp_path / "mineru-source"
    directory.mkdir()
    (directory / "full.md").write_text("安全管理\n班组监管设备", encoding="utf-8")
    (directory / "sample_content_list.json").write_text(json.dumps([
        {"type": "text", "text": "安全管理", "text_level": 1, "page_idx": 0},
        {"type": "text", "text": "班组监管设备", "page_idx": 0},
    ], ensure_ascii=False), encoding="utf-8")
    return directory


@pytest.fixture
def pdf_environment(monkeypatch, tmp_path):
    directory = _pdf_fixture(tmp_path)
    monkeypatch.setattr(worker, "_resolve_mineru_dir", lambda *args: str(directory))
    # Runtime index enhancement has independent tests; its failure must not stop KG.
    monkeypatch.setattr(worker, "_build_runtime_indexes_from_mineru_output", lambda *args, **kwargs: None)
    monkeypatch.setattr(worker, "_run_langextract", lambda documents: [])
    import services.event_bus as event_bus
    import services.audit_logger as audit_logger
    events = []
    audits = []
    monkeypatch.setattr(event_bus, "get_event_bus", lambda: SimpleNamespace(publish=lambda event, data: events.append((event, data))))
    monkeypatch.setattr(audit_logger, "get_audit_logger", lambda: SimpleNamespace(log=lambda event, **data: audits.append((event, data))))
    return directory, events, audits


@pytest.mark.parametrize("broken", ["runtime", "publish", "audit", None])
def test_pdf_optional_failure_still_returns_persisted_graph(monkeypatch, tmp_path, pdf_environment, broken):
    directory, events, audits = pdf_environment

    def unavailable(*args, **kwargs):
        raise OSError("optional subsystem unavailable")

    if broken == "runtime":
        monkeypatch.setattr(worker, "_build_runtime_indexes_from_mineru_output", unavailable)
    elif broken == "publish":
        import services.event_bus as event_bus
        monkeypatch.setattr(event_bus, "get_event_bus", unavailable)
    elif broken == "audit":
        import services.audit_logger as audit_logger
        monkeypatch.setattr(audit_logger, "get_audit_logger", unavailable)
    transitions = []
    result = worker.run_real_pipeline("pdf-task", "source.pdf", "手册.pdf", lambda *args, **kwargs: transitions.append((args, kwargs)))
    assert result["document_id"] == "doc_pdf-task"
    assert result["kg_nodes"] == result["kg_edges"] == result["grounded_entities"] == 0
    assert result["partial"] is False and result["warnings"] == []
    assert json.loads(Path(result["kg_path"], "nodes.json").read_text(encoding="utf-8"))["nodes"] == []
    assert json.loads(Path(result["kg_path"], "edges.json").read_text(encoding="utf-8"))["edges"] == []
    assert transitions[-1][0][0] == "kg_build"
    assert transitions[-1][0][3] == {"kg_build": "done"}
    assert bool(events) == (broken != "publish")
    assert bool(audits) == (broken != "audit")
    if events:
        assert events[0][0] == "doc.ingested"
        assert events[0][1]["filename"] == "手册.pdf"
        assert events[0][1]["doc_id"] == "doc_pdf-task"


@pytest.mark.parametrize("stage,error_type,code", [
    ("bridge", StageBridgeError, ErrorCode.PIPELINE_BRIDGE_PARSE),
    ("extraction", StageExtractionError, ErrorCode.PIPELINE_EXTRACTION_UNKNOWN),
    ("grounding", StageGroundingError, ErrorCode.PIPELINE_GROUNDING_ERROR),
    ("kg_build", StageKGBuildError, ErrorCode.PIPELINE_KG_BUILD_IO),
])
def test_pdf_stage_failures_report_terminal_error_and_preserve_cause(monkeypatch, tmp_path, pdf_environment, stage, error_type, code):
    directory, events, audits = pdf_environment
    root_cause = OSError("stage fixture failure")

    def fail(*args, **kwargs):
        raise root_cause

    if stage == "bridge":
        (directory / "full.md").unlink()
    else:
        monkeypatch.setattr(worker, {"extraction": "_run_langextract", "grounding": "resolve_grounding", "kg_build": "serialize_kg"}[stage], fail)
    transitions = []
    with pytest.raises(error_type) as captured:
        worker.run_real_pipeline("error-task", "source.pdf", "手册.pdf", lambda *args, **kwargs: transitions.append((args, kwargs)))
    assert captured.value.code == code
    assert captured.value.context.task_id == "error-task"
    assert captured.value.stage == PipelineStage(stage)
    if stage != "bridge":
        assert captured.value.__cause__ is root_cause
    else:
        assert isinstance(captured.value.__cause__, FileNotFoundError)
    args, kwargs = transitions[-1]
    assert args[0] == stage and args[2] == 0
    assert kwargs == {"error": captured.value.message, "error_code": code.value}
    assert events == audits == []


def test_local_parse_failure_reports_parse_failed_without_any_model_call(monkeypatch, tmp_path):
    transitions = []
    source = tmp_path / "corrupt.docx"
    source.write_bytes(b"not a zip archive")
    result = worker.run_real_pipeline("corrupt-task", str(source), source.name, lambda *args, **kwargs: transitions.append((args, kwargs)))
    assert result is None
    assert transitions[0][0][3] == {"parse": "running"}
    assert transitions[-1][0][0] == "parse"
    assert transitions[-1][0][3] == {"parse": "failed"}
    assert transitions[-1][1]["error"]
    assert not (tmp_path / "output").exists()


def test_mineru_v2_projects_titles_lists_tables_and_ignores_unknown_items(tmp_path):
    path = tmp_path / "sample_content_list_v2.json"
    path.write_text(json.dumps([
        "not a page",
        [None,
         {"type": "title", "content": {"title_content": [{"content": "附件"}, None, {}], "level": 2}},
         {"type": "paragraph", "page_idx": 4, "content": {"paragraph_content": [{"content": " scenario_summary.csv "}]}},
         {"type": "list", "page_idx": 4, "content": {"list_items": [None, {"item_content": [{"content": "W1S1 情形甲"}]}, {"item_content": [{"content": "w2s3 情形乙"}]}]}},
         {"type": "table", "page_idx": 5, "content": {"table_caption": [{"content": "统计"}], "html": "<tr><td>名</td></tr><tr><td>值</td></tr>"}},
         {"type": "title", "content": {}},
         {"type": "paragraph", "content": {}},
         {"type": "image", "content": {}},
        ],
    ], ensure_ascii=False), encoding="utf-8")
    sections, offset = worker._mineru_v2_support_sections(str(path), current_title="手册", current_level=1, char_offset=23)
    assert len(sections) == 3
    assert [section.title for section in sections] == ["附件"] * 3
    assert [section.level for section in sections] == [2] * 3
    assert [section.page_idx for section in sections] == [4, 4, 5]
    assert sections[1].text == "W1S1 情形甲\nw2s3 情形乙"
    assert sections[2].content_type == "table"
    assert sections[2].table_json == {"headers": ["名"], "rows": [["值"]], "caption": "统计"}
    assert sections[0].char_start == 23
    assert offset == sections[-1].char_end == 23 + sum(len(section.text) for section in sections)
    summary = worker._support_artifact_runtime_sections(sections, base_title="附件")
    assert len(summary) == 1
    assert [row[0] for row in summary[0].table_json["rows"]] == ["W1S1", "W2S3"]
    assert summary[0].char_start == sections[1].char_start
    assert summary[0].char_end == sections[1].char_end


def test_support_artifact_marker_stops_at_non_scenario_text_and_flushes_on_new_marker():
    def section(text, start):
        return Section(title="", level=1, text=text, content_type="text", page_idx=1, char_start=start, char_end=start + len(text))
    sections = [section("W0S1 before marker", 0), section("", 10),
                section("scenario_summary.csv", 20), section("W1S1 first", 40),
                section("scenario_summary.csv", 60), section("W1S2 second", 80),
                section("end of list", 100), section("W9S9 should stay ordinary text", 120)]
    tables = worker._support_artifact_runtime_sections(sections, base_title="")
    assert len(tables) == 2
    assert [table.table_json["rows"] for table in tables] == [[["W1S1", "W1S1 first"]], [["W1S2", "W1S2 second"]]]
    assert all(table.title == "support_artifact" for table in tables)
    assert [(table.char_start, table.char_end) for table in tables] == [(40, 50), (80, 91)]


def test_mineru_projection_loads_v2_and_sidecar_files_without_duplicate_legacy(tmp_path):
    directory = _pdf_fixture(tmp_path)
    (directory / "sample_content_list_v2.json").write_text(json.dumps([[{
        "type": "paragraph", "page_idx": 2,
        "content": {"paragraph_content": [{"content": "note.txt"}]},
    }]]), encoding="utf-8")
    (directory / "note.txt").write_text("独立附件内容", encoding="utf-8")
    sections = worker._mineru_output_to_sections(str(directory), "手册.pdf")
    assert [section.text for section in sections] == ["安全管理", "班组监管设备", "note.txt", "note.txt\n独立附件内容"]
    assert [section.title for section in sections] == ["安全管理"] * 4
    assert sections[2].page_idx == 2
    assert sections[3].page_idx is None
    assert all(right.char_start == left.char_end for left, right in zip(sections, sections[1:]))


def test_missing_mineru_outputs_fail_before_sidecar_discovery(tmp_path):
    with pytest.raises(FileNotFoundError, match="MinerU output missing"):
        worker._mineru_output_to_sections(str(tmp_path), "missing.pdf")


def test_references_are_ordered_case_insensitive_and_missing_sources_are_empty(tmp_path):
    refs = worker._extract_support_artifact_references([
        SimpleNamespace(text="Case.CSV notes.md"), SimpleNamespace(text="case.csv next.json"),
        SimpleNamespace(text=None),
    ])
    assert refs == ["Case.CSV", "notes.md", "next.json"]
    assert worker._discover_support_artifact_paths(str(tmp_path / "missing"), source_path=None, references=refs) == []


@pytest.mark.parametrize("exception_name", ["Timeout", "ReadTimeout", "ConnectTimeout"])
def test_requests_timeout_classification_preserves_timeout_reason(exception_name):
    import requests
    error = classify_external_error(getattr(requests.exceptions, exception_name)("deadline"), stage=PipelineStage.EXTRACTION, provider="ollama", task_id="classified-task")
    assert error.code == ErrorCode.PIPELINE_EXTRACTION_TIMEOUT
    assert error.context.retryable is True
    assert error.context.task_id == "classified-task"


@pytest.mark.parametrize("status,expected,retryable", [
    (401, ErrorCode.PIPELINE_EXTRACTION_AUTH, False),
    (403, ErrorCode.PIPELINE_EXTRACTION_AUTH, False),
    (429, ErrorCode.PIPELINE_EXTRACTION_RATE_LIMIT, True),
    (503, ErrorCode.PIPELINE_EXTRACTION_SERVER, True),
])
def test_requests_http_error_classification_keeps_response_status(status, expected, retryable):
    import requests
    response = requests.Response()
    response.status_code = status
    exception = requests.HTTPError("provider status", response=response)
    error = classify_external_error(exception, stage=PipelineStage.EXTRACTION, provider="ollama")
    assert error.code == expected
    assert error.context.http_status == status
    assert error.context.retryable is retryable


def test_requests_json_decode_error_classification_keeps_parse_reason():
    import requests
    exception = requests.exceptions.JSONDecodeError("invalid JSON", "?", 0)
    error = classify_external_error(exception, stage=PipelineStage.EXTRACTION, provider="ollama")
    assert error.code == ErrorCode.PIPELINE_EXTRACTION_PARSE
    assert error.context.retryable is False


@pytest.mark.parametrize("exception_name", ["PoolTimeout", "WriteTimeout"])
def test_httpx_timeout_subclasses_keep_timeout_reason(exception_name):
    import httpx
    error = classify_external_error(getattr(httpx, exception_name)("deadline"), stage=PipelineStage.EXTRACTION, provider="ollama")
    assert error.code == ErrorCode.PIPELINE_EXTRACTION_TIMEOUT
    assert error.context.retryable is True


def test_provider_business_error_has_priority_and_unknown_oserror_still_falls_back():
    import requests
    # The recognized body code must be handled before exception inheritance.
    business_error = classify_external_error(requests.Timeout("deadline"), stage=PipelineStage.MINERU, provider="mineru", body_code=-60023)
    assert business_error.context.provider_code == -60023
    assert business_error.code == ErrorCode.PIPELINE_MINERU_IO
    fallback = classify_external_error(OSError("connection failed"), stage=PipelineStage.EXTRACTION, provider="ollama")
    assert fallback.code == ErrorCode.PIPELINE_EXTRACTION_SERVER
    assert fallback.context.retryable is True


@pytest.mark.parametrize("mode", ["success", "partial", "empty", "embed-missing", "embed-failed", "index-failed", "index-absent"])
def test_docx_pipeline_uses_real_chunks_and_graph_with_explicit_enhancement_degradation(monkeypatch, tmp_path, pdf_environment, mode):
    from docx import Document
    import index_registry
    import requests
    document = Document()
    document.add_heading("设备管理", level=1)
    for paragraph in ["班组监管设备并开展专项检查。", "锅炉压力必须维持规定范围。", "消防通道禁止堆放易燃杂物。", "年度预算由财务部门统一审批。", "工作人员佩戴防护眼镜进入车间。", "应急演练结束后形成总结报告。"]:
        document.add_paragraph(paragraph)
    source = tmp_path / "设备.docx"
    document.save(source)
    embedded = []
    stored = []
    indexes = {}

    class EmbeddingBoundary:
        def __init__(self):
            if mode == "embed-missing":
                raise ImportError("embedding boundary unavailable")

        def embed_texts(self, texts):
            if mode == "embed-failed":
                raise RuntimeError("embedding boundary failed")
            embedded.extend(texts)
            return [[float(index), 1.0] for index in range(len(texts))]

    class VectorBoundary:
        def add_chunks(self, chunks, embeddings):
            stored.append((chunks, embeddings))

    # No model import/download or persistent Chroma connection is involved.
    monkeypatch.setitem(sys.modules, "embedding", SimpleNamespace(EmbeddingClient=EmbeddingBoundary))
    monkeypatch.setitem(sys.modules, "vector_store", SimpleNamespace(VectorStore=VectorBoundary))
    monkeypatch.setattr(index_registry, "set_bm25_index", lambda index: indexes.update(bm25=index))
    monkeypatch.setattr(index_registry, "set_table_index", lambda index: indexes.update(table=index))
    if mode == "index-failed":
        def fail_index(chunks):
            raise RuntimeError("runtime index unavailable")
        monkeypatch.setattr(index_registry, "build_bm25_from_chunks", fail_index)
    elif mode == "index-absent":
        monkeypatch.setattr(index_registry, "build_bm25_from_chunks", lambda chunks: None)
        monkeypatch.setattr(index_registry, "build_table_index_from_chunks", lambda chunks: None)
    responses = ([requests.Timeout("first batch expired"), _graph_payload()] if mode == "partial" else
                 [{"entities": []}] * 2 if mode == "empty" else [_graph_payload()] * 2)
    calls = _provider(monkeypatch, responses)
    transitions = []
    result = worker.run_real_pipeline("docx-task", str(source), source.name, lambda *args, **kwargs: transitions.append((args, kwargs)))
    assert result["document_id"] == "doc_docx-task"
    assert result["chunks"] == result["pages_parsed"] == 6
    assert result["kg_nodes"] == (0 if mode == "empty" else 2)
    assert result["kg_edges"] == (0 if mode == "empty" else 1)
    assert result["partial"] == (mode == "partial")
    assert len(calls) == 2
    assert result["warnings"] == (["1/2 批次抽取失败"] if mode == "partial" else [])
    degradation = {args[0]: (args, kwargs) for args, kwargs in transitions if kwargs.get("degraded")}
    expected_stage = {"partial": "extraction", "empty": "extraction", "embed-missing": "embed", "embed-failed": "embed", "index-failed": "index"}.get(mode)
    assert set(degradation) == ({expected_stage} if expected_stage else set())
    if expected_stage:
        assert degradation[expected_stage][1]["error"]
        assert degradation[expected_stage][1]["error_code"]
    if mode not in {"embed-missing", "embed-failed"}:
        chunks, vectors = stored[0]
        assert embedded == [chunk.text for chunk in chunks]
        assert len(vectors) == len(chunks) == 6
        assert all(chunk.metadata.source_file == source.name and chunk.metadata.source_format == "docx" for chunk in chunks)
    else:
        assert stored == []
    if mode not in {"index-failed", "index-absent"}:
        assert len(indexes["bm25"][1]) == 6
        assert indexes["table"].count() == 0
        indexes["table"].conn.close()
    else:
        assert indexes == {}
    _, events, audits = pdf_environment
    assert events[0][1]["pages_parsed"] == 6
    assert audits[0][1]["chunks"] == 6
    if mode != "empty":
        assert json.loads(Path(result["kg_path"], "nodes.json").read_text(encoding="utf-8"))["nodes"][0]["label"] == "班组"


@pytest.mark.parametrize("list_result", [True, False])
def test_langextract_builds_real_examples_and_model_config_without_provider_call(monkeypatch, tmp_path, list_result):
    # The repository root also contains the editable project's outer directory.
    # Resolve the actual installed package, as a backend-cwd launch does, rather
    # than the outer namespace package created by a repository-root pytest run.
    import importlib
    monkeypatch.syspath_prepend(str(Path(__file__).resolve().parents[2] / "langextract"))
    monkeypatch.delitem(sys.modules, "langextract", raising=False)
    langextract = importlib.import_module("langextract")
    config_dir = tmp_path / "config"
    config_dir.mkdir()
    (config_dir / "extraction_examples.yaml").write_text(
        "examples:\n  - text: 班组负责设备\n    extractions:\n      - extraction_class: department\n        extraction_text: 班组\n        attributes:\n          dept_name: 班组\n", encoding="utf-8")
    monkeypatch.setenv("LANGEXTRACT_MODEL_ID", "test-boundary-model")
    monkeypatch.setenv("LANGEXTRACT_API_KEY", "test-only-placeholder")
    monkeypatch.setenv("LANGEXTRACT_BASE_URL", "https://example.invalid/v1")
    response = SimpleNamespace(document_id="response")
    captured = {}

    def extract(**kwargs):
        captured.update(kwargs)
        return [response] if list_result else response

    monkeypatch.setattr(langextract, "extract", extract)
    documents = [SimpleNamespace(document_id="input", text="班组负责设备")]
    result = worker._run_langextract(documents)
    assert result == [response]
    assert captured["text_or_documents"] is documents
    assert captured["examples"][0].text == "班组负责设备"
    assert captured["examples"][0].extractions[0].extraction_text == "班组"
    assert captured["examples"][0].extractions[0].attributes == {"dept_name": "班组"}
    assert captured["config"].model_id == "test-boundary-model"
    assert captured["config"].provider == "openai"
    assert captured["config"].provider_kwargs == {"api_key": "test-only-placeholder", "base_url": "https://example.invalid/v1"}
    assert captured["batch_length"] == 5 and captured["max_workers"] == 3
    assert captured["show_progress"] is False and captured["temperature"] == 0


@pytest.mark.parametrize("failed", [False, True])
def test_public_upload_uses_unique_basename_and_returns_none_on_boundary_failure(monkeypatch, tmp_path, failed):
    source = tmp_path / "设备.pdf"
    source.write_bytes(b"local upload boundary fixture")
    calls = []

    def upload(path, name):
        assert Path(path).read_bytes() == b"local upload boundary fixture"
        calls.append(name)
        if failed:
            raise OSError("upload boundary unavailable")
        return "https://example.invalid/设备.pdf"

    monkeypatch.setitem(sys.modules, "oss_uploader", SimpleNamespace(upload_to_oss=upload))
    result = worker._upload_to_public(str(source))
    assert result == (None if failed else "https://example.invalid/设备.pdf")
    assert len(calls) == 1
    prefix, basename = calls[0].split("_", 1)
    assert basename == source.name and len(prefix) == 8
    assert all(character in "0123456789abcdef" for character in prefix)
