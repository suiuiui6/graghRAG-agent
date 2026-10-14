"""Real-format temporary fixtures and document processing behavior contracts."""
import json

import pytest
from docx import Document
from openpyxl import Workbook

from doc_processor import (DocumentProcessor, Section, chunk_sections, clean_text,
                           load_chunks_jsonl, load_project_id_mapping, parse_document)


def test_docx_headings_blank_paragraphs_and_structured_table(tmp_path):
    document = Document()
    document.add_heading("Overview", level=2)
    document.add_paragraph("")
    document.add_paragraph("A bounded project description.")
    table = document.add_table(rows=2, cols=2)
    for row, values in zip(table.rows, [("Item", "Value"), ("Budget", "42")]):
        for cell, value in zip(row.cells, values):
            cell.text = value
    path = tmp_path / "report.docx"
    document.save(path)
    sections = parse_document(path)
    assert len(sections) == 2
    assert sections[0].title == "Overview" and sections[0].level == 2
    assert sections[0].text == "A bounded project description."
    assert sections[1].content_type == "table"
    assert sections[1].table_json == {"headers": ["Item", "Value"], "rows": [["Budget", "42"]]}
    assert sections[1].char_start == sections[0].char_end


def test_workbook_sheets_numeric_cells_empty_rows_and_header_only_sheet(tmp_path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Name", "Value"])
    sheet.append([None, None])
    sheet.append(["Revenue", 42])
    sheet.append(["Unreported", None])
    workbook.create_sheet("Empty").append(["Name"])
    path = tmp_path / "report.xlsx"
    workbook.save(path)
    workbook.close()
    sections = parse_document(path)
    assert len(sections) == 1
    assert sections[0].title == "Metrics"
    assert sections[0].table_json["rows"] == [["Revenue", "42"], ["Unreported", ""]]


@pytest.mark.parametrize("level", [1, 2, 3, 999])
def test_document_process_save_reload_keeps_text_identity_and_project(level, tmp_path):
    path = tmp_path / "report.md"
    path.write_text("# Empty heading\n## Overview\nA bounded statement.\n###### Details\nAnother fact.", encoding="utf-8")
    processor = DocumentProcessor(chunk_strategy=level, dedup=False,
                                  project_id_mapping={"report.md": "fixture-project"})
    chunks = processor.process(path)
    assert [chunk.text for chunk in chunks] == ["A bounded statement.", "Another fact."]
    assert [chunk.chunk_id for chunk in chunks] == ["c00001", "c00002"]
    assert all(chunk.metadata.project_id == "fixture-project" for chunk in chunks)
    assert all(chunk.metadata.source_format == "md" for chunk in chunks)
    assert chunks[1].metadata.section_path == ["Details"]
    target = tmp_path / "saved" / "chunks.jsonl"
    processor.save(chunks, str(target))
    restored = load_chunks_jsonl(str(target))
    assert restored == chunks


def test_plain_markdown_and_unsupported_format(tmp_path):
    path = tmp_path / "plain.md"
    path.write_text("A plain document.", encoding="utf-8")
    section, = parse_document(path)
    assert section.title == "plain.md" and section.text == "A plain document."
    with pytest.raises(ValueError, match="Unsupported format"):
        parse_document(tmp_path / "unsupported.exe")
    with pytest.raises(FileNotFoundError, match="Incomplete MinerU"):
        parse_document(tmp_path)


def test_directory_continues_after_corrupt_document_and_deduplicates(tmp_path):
    (tmp_path / "a.md").write_text("Same content for both files.", encoding="utf-8")
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "b.md").write_text("Same content for both files.", encoding="utf-8")
    (nested / "broken.docx").write_bytes(b"not a ZIP document")
    (nested / "ignore.txt").write_text("Excluded by supported formats.", encoding="utf-8")
    chunks = DocumentProcessor(project_id_mapping={}).process_directory(str(tmp_path))
    assert len(chunks) == 1
    assert chunks[0].text == "Same content for both files."


def test_mapping_missing_invalid_and_non_string_values(tmp_path):
    path = tmp_path / "mapping.json"
    assert load_project_id_mapping(str(path)) == {}
    path.write_text("invalid json", encoding="utf-8")
    assert load_project_id_mapping(str(path)) == {}
    path.write_text(json.dumps({"a.md": "project", "invalid": 3}), encoding="utf-8")
    assert load_project_id_mapping(str(path)) == {"a.md": "project"}
    assert DocumentProcessor(project_id_mapping=str(path))._mapping == {"a.md": "project"}


@pytest.mark.parametrize("value", [[], ["project"], "project", 42, None])
def test_mapping_rejects_valid_json_with_wrong_top_level_type(tmp_path, value):
    path = tmp_path / "mapping.json"
    path.write_text(json.dumps(value), encoding="utf-8")
    assert load_project_id_mapping(str(path)) == {}


def test_fixed_chunks_overlap_and_offsets_match_source():
    text = "abcdefghijklmnopqrstuvwxyz"
    section = Section("Text", 1, text, "text", char_start=100, char_end=126)
    chunks = chunk_sections([section], "file.md", "md", chunk_size=4, overlap=1, level=1, mapping={})
    assert len(chunks) > 1
    for chunk in chunks:
        start, end = chunk.metadata.char_start - 100, chunk.metadata.char_end - 100
        assert chunk.text == text[start:end]
    assert chunks[-1].metadata.char_end == 126
    for first, second in zip(chunks, chunks[1:]):
        assert second.metadata.char_start < first.metadata.char_end


def test_clean_text_removes_only_standalone_page_furniture():
    assert clean_text("Page 1 of 2\n\nConfidential\n\nActual statement 42.\n\n123") == "Actual statement 42."
    assert "Page 1 of 2 is a quoted label" in clean_text("Page 1 of 2 is a quoted label")
