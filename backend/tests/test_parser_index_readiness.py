"""Real document artifacts and index recovery behavior in isolated directories."""
import builtins
import json
import sys
from types import SimpleNamespace

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from openpyxl import Workbook

import doc_processor as dp
import index_registry as ir


def test_mineru_real_bridge_and_pdf_cache_selection(tmp_path, monkeypatch):
    monkeypatch.setattr(dp, '__file__', str(tmp_path / 'doc_processor.py'))
    with pytest.raises(FileNotFoundError, match='No MinerU output'):
        dp.parse_document(tmp_path / 'missing.pdf')
    root = tmp_path / 'output' / 'mineru'
    root.mkdir(parents=True)
    (root / 'zzz-file').write_text('not a directory')
    (root / 'zz-empty').mkdir()
    with pytest.raises(FileNotFoundError, match='No MinerU output'):
        dp.parse_document(tmp_path / 'missing.pdf')
    cache = root / 'a-valid'
    cache.mkdir()
    text = 'Introduction\nPump output is 42.\n'
    (cache / 'full.md').write_text(text, encoding='utf-8')
    items = [{'type': 'text', 'text': 'Introduction', 'text_level': 1, 'page_idx': 0},
             {'type': 'text', 'text': 'Pump output is 42.', 'page_idx': 1}]
    (cache / 'pump_content_list.json').write_text(json.dumps(items), encoding='utf-8')
    (cache / 'pump_v2_content_list.json').write_text('invalid', encoding='utf-8')
    direct = dp.parse_document(cache)
    cached = dp.parse_document(tmp_path / 'pump.pdf')
    assert direct == cached and len(direct) == 1
    assert direct[0].title == 'Introduction' and '42' in direct[0].text
    assert direct[0].page_idx == 0
    (cache / 'pump_content_list.json').write_text('invalid', encoding='utf-8')
    with pytest.raises(json.JSONDecodeError):
        dp.parse_document(cache)


def test_docx_non_numeric_heading_and_absent_xml_lookups(tmp_path):
    doc = Document()
    doc.styles.add_style('Heading Custom', WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph('Fallback title', 'Heading Custom')
    doc.add_paragraph('Useful body')
    empty_table = doc.add_table(rows=0, cols=1)
    assert dp._extract_table_data_from_docx(empty_table) == {'headers': ['col_0'], 'rows': []}
    assert dp._find_paragraph_in_doc(doc, OxmlElement('w:p')) is None
    assert dp._find_table_in_doc(doc, OxmlElement('w:tbl')) is None
    target = tmp_path / 'heading.docx'
    doc.save(target)
    sections = dp.parse_document(target)
    assert sections[0].title == 'Fallback title' and sections[0].level == 1
    assert sections[1].content_type == 'table' and sections[1].text == ''


@pytest.mark.parametrize('text,expected', [('', []), ('Unmapped content', [None])])
def test_mineru_empty_or_unmapped_position_has_no_fabricated_page(tmp_path, text, expected):
    (tmp_path / 'full.md').write_text(text, encoding='utf-8')
    (tmp_path / 'doc_content_list.json').write_text('[]', encoding='utf-8')
    assert [section.page_idx for section in dp.parse_document(tmp_path)] == expected


@pytest.mark.parametrize('empty_prefix,headers', [(1, ['name', 'value']), (2, ['col_0', 'col_1'])])
def test_xlsx_header_recovery_after_empty_prefix(tmp_path, empty_prefix, headers):
    wb = Workbook()
    ws = wb.active
    for _ in range(empty_prefix):
        ws.append([None, None])
    ws.append(['name', 'value'])
    ws.append(['budget', 12])
    target = tmp_path / 'prefix.xlsx'
    wb.save(target)
    wb.close()
    section, = dp.parse_document(target)
    assert section.table_json['headers'] == headers
    assert section.table_json['rows'][-1] == ['budget', '12']


def test_long_semantic_chunks_sentence_boundaries_and_project_sources(tmp_path):
    text = ('设备运行正常。' * 1600)
    chunks = dp.chunk_sections([dp.Section('Details', 2, text, 'text', char_end=len(text))],
                               'report.md', 'md', mapping={})
    assert len(chunks) > 1 and ''.join(c.text for c in chunks) == text
    assert all(len(c.text) <= 8000 for c in chunks)
    assert chunks[1].metadata.section_path == ['Details', 'part_2']
    assert dp._split_long_text('abc', 4) == ['abc']
    assert dp._split_long_text('abc     ', 3) == ['abc']
    assert dp.extract_project_id('unmapped.md', 'Report PJ-2026-042', '', {'x': 'other'}) == 'PJ-2026-042'
    assert dp.extract_project_id('unmapped.md', '', '项目编号: ABC_42') == 'ABC_42'
    assert dp.extract_project_id('项目编号: X42.md', '', '') == 'X42'
    assert dp._table_to_text({'rows': [[1, 2]]}) == '1 | 2'
    assert dp._jaccard_similarity(set(), {'a'}) == 0
    assert dp.load_project_id_mapping(str(tmp_path / 'absent')) == {}
    mapping = tmp_path / 'mapping.json'
    mapping.write_text('[]')
    assert dp.load_project_id_mapping(str(mapping)) == {}
    target = tmp_path / 'chunks.jsonl'
    dp.save_chunks_jsonl([dp.Chunk('plain', 'text', 'text')], str(target))
    with target.open('a', encoding='utf-8') as output:
        output.write('\n')
    restored, = dp.load_chunks_jsonl(str(target))
    assert restored.text == 'text' and restored.metadata.source_file == ''
    assert dp.DocumentProcessor(dedup=False, project_id_mapping={}).process_directory(str(tmp_path / 'missing')) == []


@pytest.fixture
def indexes(monkeypatch, tmp_path):
    monkeypatch.setattr(ir, '_bm25_index', None)
    monkeypatch.setattr(ir, '_table_index', None)
    monkeypatch.setattr(ir, 'BM25_ARTIFACT_PATH', tmp_path / 'bm25.json')
    monkeypatch.setattr(ir, 'TABLE_ARTIFACT_PATH', tmp_path / 'table.json')
    return ir


def test_real_indexes_snapshots_restore_and_rank(indexes):
    chunks = [dp.Chunk('a', 'pump pressure safety', 'text'), dp.Chunk('b', 'operator budget', 'text'),
              dp.Chunk('t', 'metric 42', 'table', table_json={'headers': ['metric', 'value'], 'rows': [['pressure', '42']]})]
    bm = indexes.build_bm25_from_chunks(chunks)
    assert bm[0].get_scores(['pump'])[0] > bm[0].get_scores(['pump'])[1]
    table = indexes.build_table_index_from_chunks(chunks)
    assert table is not None and table.query(top_k=5)
    indexes.set_bm25_index(bm)
    indexes.set_table_index(table)
    indexes.set_bm25_index(None, persist=False)
    indexes.set_table_index(None, persist=False)
    state = indexes.restore_runtime_indexes()
    assert state['bm25']['restored'] and state['table']['restored']
    assert indexes.get_bm25_index()[1][0].text == chunks[0].text
    restored = indexes.get_table_index()
    assert restored.query(conditions=['value > 0']) == []
    assert restored.get_stats()['restored'] and restored.get_stats()['total_rows'] > 0
    assert indexes.get_runtime_recovery_status()['table']['loaded']
    table.conn.close()


@pytest.mark.parametrize('payload', [[], {'chunks': []}, {'chunks': 'bad'}, {'chunks': [23]}])
def test_invalid_bm25_does_not_block_valid_table(indexes, payload):
    indexes.persist_bm25_artifact(indexes.BM25_ARTIFACT_PATH, payload)
    indexes.persist_table_artifact(indexes.TABLE_ARTIFACT_PATH, {'rows': [{'sheet_name': 'One'}, {}, {'sheet_name': 'One'}]})
    result = indexes.restore_runtime_indexes()
    assert not result['bm25']['restored'] and result['table']['restored']
    assert indexes.get_table_index().get_stats() == {'total_rows': 3, 'unique_sheets': 1, 'wide_tables': 0, 'restored': True}


@pytest.mark.parametrize('payload', [[], {'rows': None}, {'rows': 'bad'}])
def test_invalid_table_with_valid_legacy_bm25(indexes, payload):
    indexes.persist_bm25_artifact(indexes.BM25_ARTIFACT_PATH, {'chunks': ['pump safety', {'text': 'budget cost'}]})
    indexes.persist_table_artifact(indexes.TABLE_ARTIFACT_PATH, payload)
    result = indexes.restore_runtime_indexes()
    assert result['bm25']['restored'] and not result['table']['restored']
    assert indexes.get_bm25_index()[1][0].chunk_id == 'restored_bm25_0'


def test_index_empty_corrupt_artifacts_and_legacy_query_signatures(indexes):
    assert indexes.snapshot_bm25_index() is None
    assert indexes.snapshot_table_index() is None
    assert indexes.build_bm25_from_chunks([]) is None
    assert indexes.restore_runtime_indexes()['bm25']['reason'] == 'not_found'
    indexes.BM25_ARTIFACT_PATH.write_text('{bad')
    assert indexes.restore_bm25_artifact(indexes.BM25_ARTIFACT_PATH) is None
    # A directory is a real unreadable artifact for Path.read_text on Windows/Linux.
    indexes.TABLE_ARTIFACT_PATH.mkdir()
    assert indexes.restore_table_artifact(indexes.TABLE_ARTIFACT_PATH) is None
    assert indexes._serialize_bm25_index(('bad',)) is None
    class LegacyTable:
        def query(self, conditions, count):
            assert conditions is None and isinstance(count, int) and count > 0
            return [{'sheet_name': 'old', 'value': 42}]
    assert indexes._serialize_table_index(LegacyTable()) == {'rows': [{'sheet_name': 'old', 'value': 42}]}
    class FailedTable:
        def query(self, *args, **kwargs):
            raise OSError('broken database')
    assert indexes._serialize_table_index(FailedTable()) is None
    assert indexes._serialize_chunk('legacy')['text'] == 'legacy'
    assert indexes._tokenize_for_bm25('   ') == []
    assert indexes._fallback_tokenize_for_bm25('pump  中文 x') == ['pump', '中文', '文', 'x']


def test_optional_parser_dependencies_report_actionable_errors(monkeypatch, tmp_path):
    original = builtins.__import__
    def unavailable(name, *args, **kwargs):
        if name in {'docx', 'openpyxl'}:
            raise ImportError('dependency unavailable')
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', unavailable)
    with pytest.raises(ImportError, match='python-docx is required'):
        dp.parse_document(tmp_path / 'file.docx')
    with pytest.raises(ImportError, match='openpyxl is required'):
        dp.parse_document(tmp_path / 'file.xlsx')


def test_mapping_default_missing_and_fixed_chunks_ignore_internal_whitespace(tmp_path, monkeypatch):
    monkeypatch.setattr(dp, '__file__', str(tmp_path / 'doc_processor.py'))
    assert dp.DocumentProcessor()._mapping == {}
    chunks = dp.chunk_sections([dp.Section('Spaces', 1, 'abcd        efgh', 'text')],
                               'file.md', 'md', chunk_size=1, overlap=0, level=1, mapping={})
    assert [c.text for c in chunks] == ['abcd', 'efgh']
    assert dp.extract_project_id('file.md', 'No project match', 'PJ-2026-001', {}) == 'PJ-2026-001'
    assert dp.extract_project_id('file.md', '', 'No match', {}) is None
    originals = [dp.Chunk('a', 'alpha beta', 'text'), dp.Chunk('b', 'alpha beta alpha beta', 'text'),
                 dp.Chunk('c', 'different words', 'text')]
    assert [c.chunk_id for c in dp.deduplicate_chunks(originals)] == ['b', 'c']


def test_index_missing_dependencies_and_invalid_artifact_shapes(indexes, monkeypatch):
    original = builtins.__import__
    def unavailable(name, *args, **kwargs):
        if name in {'rank_bm25', 'table_index', 'jieba', 'routes.rag_query'}:
            raise ImportError('optional dependency unavailable')
        return original(name, *args, **kwargs)
    monkeypatch.setattr(builtins, '__import__', unavailable)
    assert indexes.build_bm25_from_chunks([dp.Chunk('a', 'pump', 'text')]) is None
    assert indexes.build_table_index_from_chunks([]) is None
    assert indexes._tokenize_for_bm25('Pump 中文') == ['pump', '中文', '文']
    # Persist to real directory paths: writes fail, while in-memory state still advances.
    indexes.BM25_ARTIFACT_PATH.mkdir()
    indexes.TABLE_ARTIFACT_PATH.mkdir()
    bm = (object(), [dp.Chunk('a', 'pump', 'text')])
    indexes.set_bm25_index(bm)
    indexes.set_table_index(indexes.RestoredTableIndex([{'sheet_name': 's'}]))
    assert indexes.get_bm25_index() is bm and indexes.get_table_index().query()[0]['sheet_name'] == 's'
    assert indexes._deserialize_bm25_index('wrong shape') is None
    assert indexes._deserialize_table_index('wrong shape') is None


def test_index_deserialization_exceptions_are_isolated_and_legacy_query_failure(indexes):
    # Wrong per-chunk text type raises inside real BM25 tokenization; table still restores.
    indexes.persist_bm25_artifact(indexes.BM25_ARTIFACT_PATH, {'chunks': [{'text': None}]})
    indexes.persist_table_artifact(indexes.TABLE_ARTIFACT_PATH, {'rows': [{'sheet_name': 'safe'}]})
    result = indexes.restore_runtime_indexes()
    assert result['bm25']['reason'] == 'invalid_payload' and result['table']['restored']
    class FailedLegacyTable:
        def query(self, conditions, count):
            raise OSError('legacy database missing')
    assert indexes._serialize_table_index(FailedLegacyTable()) is None
    # Invalid Chunk content produces an actual ingest failure, caught by the builder.
    assert indexes.build_table_index_from_chunks([SimpleNamespace()]) is None


def test_index_status_and_optional_tokenizer_boundary(indexes, monkeypatch):
    status = indexes.get_runtime_recovery_status()
    assert status['bm25']['loaded'] is False and status['table']['artifact_exists'] is False
    # jieba is optional in the candidate lock; exercise its protocol boundary only.
    monkeypatch.setitem(sys.modules, 'jieba', SimpleNamespace(cut=lambda text: iter([' ', ' pump ', '中文'])))
    assert indexes._tokenize_for_bm25('PUMP 中文') == ['pump', '中文']
    monkeypatch.setitem(sys.modules, 'jieba', SimpleNamespace(cut=lambda text: iter([' ', ''])))
    assert indexes._tokenize_for_bm25('pump中文') == ['pump', '中文', '文']
