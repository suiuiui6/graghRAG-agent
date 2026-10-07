from pathlib import Path
import math
import textwrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\graghRAG-agent")
OUT_DIR = ROOT / "docs" / "superpowers" / "plans"
FIG_DIR = ROOT / "docs" / "superpowers" / "figures" / "graphrag_core_journal"
DOCX_PATH = OUT_DIR / "2026-06-17-graphrag-paper-core-journal-with-figures-v1.docx"


def font(size=28, bold=False):
    candidates = [
        r"C:\Windows\Fonts\msyhbd.ttc" if bold else r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        r"C:\Windows\Fonts\simsun.ttc",
        r"C:\Windows\Fonts\arial.ttf",
    ]
    for p in candidates:
        if Path(p).exists():
            return ImageFont.truetype(p, size=size)
    return ImageFont.load_default()


FONT_TITLE = font(42, True)
FONT_H = font(30, True)
FONT = font(24)
FONT_SMALL = font(20)
FONT_TINY = font(17)


def wrap_cn(text, width):
    lines = []
    cur = ""
    units = text.replace("，", "， ").replace("；", "； ").replace("。", "。 ").split()
    for u in units:
        if len(cur) + len(u) <= width:
            cur += u
        else:
            if cur:
                lines.append(cur)
            cur = u
    if cur:
        lines.append(cur)
    if not lines:
        return [text]
    out = []
    for line in lines:
        if len(line) <= width:
            out.append(line)
        else:
            out.extend(textwrap.wrap(line, width=width))
    return out


def draw_centered_text(draw, box, text, fnt, fill=(35, 45, 55), line_spacing=6, max_chars=15):
    x0, y0, x1, y1 = box
    lines = []
    for part in str(text).split("\n"):
        lines.extend(wrap_cn(part, max_chars))
    heights = []
    widths = []
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=fnt)
        widths.append(bbox[2] - bbox[0])
        heights.append(bbox[3] - bbox[1])
    total_h = sum(heights) + line_spacing * (len(lines) - 1)
    y = y0 + (y1 - y0 - total_h) / 2
    for line, w, h in zip(lines, widths, heights):
        x = x0 + (x1 - x0 - w) / 2
        draw.text((x, y), line, font=fnt, fill=fill)
        y += h + line_spacing


def rounded_box(draw, box, fill, outline=(70, 90, 110), radius=22, width=2):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, fill=(70, 90, 110), width=4):
    draw.line([start, end], fill=fill, width=width)
    x1, y1 = start
    x2, y2 = end
    angle = math.atan2(y2 - y1, x2 - x1)
    head_len = 16
    left = (x2 - head_len * math.cos(angle - math.pi / 6), y2 - head_len * math.sin(angle - math.pi / 6))
    right = (x2 - head_len * math.cos(angle + math.pi / 6), y2 - head_len * math.sin(angle + math.pi / 6))
    draw.polygon([end, left, right], fill=fill)


def base_canvas(title):
    img = Image.new("RGB", (1800, 1050), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((0, 0, 1800, 1050), fill=(250, 252, 253))
    d.text((70, 45), title, font=FONT_TITLE, fill=(20, 35, 50))
    d.line((70, 105, 1730, 105), fill=(180, 195, 210), width=3)
    return img, d


def make_fig1():
    img, d = base_canvas("图1  面向中文长文档问答的可追溯 GraphRAG 系统架构")
    boxes = [
        ((80, 210, 310, 370), "中文长文档\nPDF/报告/论文", (233, 242, 251)),
        ((380, 210, 650, 370), "文档解析与\n知识图谱证据组织", (226, 245, 237)),
        ((720, 210, 970, 370), "查询预处理\n问题改写/约束", (252, 241, 222)),
        ((1040, 210, 1310, 370), "多路径检索\n节点/关系/邻居", (238, 232, 250)),
        ((1380, 210, 1690, 370), "重排与\nfinal-context 选择", (232, 246, 250)),
        ((380, 585, 650, 745), "预算约束控制\n调用/token/延迟", (255, 239, 235)),
        ((720, 585, 970, 745), "引用校验生成\n引用来自已检索证据", (232, 246, 250)),
        ((1040, 585, 1310, 745), "拒答与\n假前提纠错", (246, 238, 231)),
        ((1380, 585, 1690, 745), "可追溯回答\n证据/边界/成本", (226, 245, 237)),
    ]
    for b, t, c in boxes:
        rounded_box(d, b, c)
        draw_centered_text(d, b, t, FONT, max_chars=10)
    for i in range(4):
        arrow(d, (boxes[i][0][2], 290), (boxes[i + 1][0][0], 290))
    arrow(d, (1535, 370), (1535, 585))
    arrow(d, (1380, 665), (1310, 665))
    arrow(d, (1040, 665), (970, 665))
    arrow(d, (720, 665), (650, 665))
    d.line((515, 585, 515, 430, 1510, 430, 1510, 370), fill=(130, 95, 80), width=4)
    d.text((585, 445), "预算控制贯穿检索扩张、证据筛选与生成阶段", font=FONT_SMALL, fill=(120, 75, 60))
    d.text((80, 920), "说明：图中 final-context 为生成前最终证据集合，是论文默认检索主结果面；path_hits 与 pre-rerank 仅用于诊断。", font=FONT_SMALL, fill=(75, 85, 95))
    out = FIG_DIR / "fig1_system_architecture.png"
    img.save(out, dpi=(220, 220))
    return out


def make_fig2():
    img, d = base_canvas("图2  证据面分层与论文写作准入规则")
    stages = [
        ("path_hits\n早期路径命中", "诊断面\n只分析证据是否曾被碰到", (238, 232, 250)),
        ("pre-rerank\n重排前候选", "诊断面\n分析候选池和丢证据阶段", (238, 232, 250)),
        ("final-context\n生成前最终证据", "主结果面\n摘要/引言/主结果表准入", (226, 245, 237)),
        ("answer-support\n回答支持", "支撑面\n检查引用真实性与 grounding", (232, 246, 250)),
        ("release gate\n工程门禁", "工程面\n说明复现状态，不能替代主结论", (255, 239, 235)),
    ]
    x = 80
    for idx, (name, desc, color) in enumerate(stages):
        b = (x, 230, x + 290, 480)
        rounded_box(d, b, color)
        draw_centered_text(d, (b[0] + 10, b[1] + 25, b[2] - 10, b[1] + 115), name, FONT_H, max_chars=12)
        draw_centered_text(d, (b[0] + 20, b[1] + 135, b[2] - 20, b[3] - 25), desc, FONT_SMALL, max_chars=12)
        if idx < len(stages) - 1:
            arrow(d, (b[2], 355), (b[2] + 55, 355))
        x += 340
    d.rectangle((130, 650, 1670, 820), fill=(245, 248, 250), outline=(170, 185, 200), width=2)
    d.text((165, 675), "写作准入：", font=FONT_H, fill=(30, 50, 70))
    rules = [
        "1. final-context 是默认 headline retrieval surface。",
        "2. path_hits 与 pre-rerank 只能进入诊断、失败分析或附录。",
        "3. historical evidence 与 current frozen evidence 必须分开。",
        "4. release gate 是工程支撑，不包装为学术核心贡献。",
    ]
    yy = 680
    for r in rules:
        yy += 34
        d.text((330, yy), r, font=FONT_SMALL, fill=(60, 70, 80))
    out = FIG_DIR / "fig2_evidence_surfaces.png"
    img.save(out, dpi=(220, 220))
    return out


def make_fig3():
    img, d = base_canvas("图3  历史快照与当前工程门禁的主要信号")
    metrics = [
        ("Recall\nmacro", 0.82, (82, 145, 190)),
        ("Precision\nmacro", 0.078, (210, 120, 90)),
        ("Keyword\ncoverage", 0.799, (90, 160, 120)),
        ("Cited in\nretrieved", 1.0, (90, 160, 120)),
        ("Gate final-context\nrecall", 0.5765, (120, 120, 190)),
        ("Gate keyword\ncoverage", 0.6893, (210, 120, 90)),
    ]
    chart = (120, 210, 1260, 790)
    d.line((chart[0], chart[3], chart[2], chart[3]), fill=(120, 135, 150), width=3)
    d.line((chart[0], chart[1], chart[0], chart[3]), fill=(120, 135, 150), width=3)
    for t in [0, .25, .5, .75, 1.0]:
        y = chart[3] - t * (chart[3] - chart[1])
        d.line((chart[0] - 8, y, chart[2], y), fill=(225, 232, 238), width=1)
        d.text((65, y - 12), f"{t:.2f}", font=FONT_TINY, fill=(80, 90, 100))
    bar_w = 115
    gap = 70
    x = chart[0] + 55
    for label, val, color in metrics:
        h = val * (chart[3] - chart[1])
        d.rounded_rectangle((x, chart[3] - h, x + bar_w, chart[3]), radius=12, fill=color)
        d.text((x - 15, chart[3] + 18), label, font=FONT_TINY, fill=(40, 50, 60))
        d.text((x + 8, chart[3] - h - 32), f"{val:.3f}".rstrip("0").rstrip("."), font=FONT_SMALL, fill=(30, 45, 60))
        x += bar_w + gap
    # release gate summary panel
    panel = (1330, 220, 1690, 790)
    rounded_box(d, panel, (255, 247, 238), outline=(210, 160, 120))
    d.text((1370, 255), "当前 release gate", font=FONT_H, fill=(110, 65, 45))
    gate_lines = [
        "Final verdict: NO-GO",
        "PASS: final-context recall",
        "PASS: cited_in_retrieved",
        "PASS: p95/p99 latency",
        "FAIL: keyword coverage",
        "FAIL: backend validation",
    ]
    yy = 320
    for line in gate_lines:
        color = (120, 70, 55) if "NO-GO" in line or "FAIL" in line else (55, 115, 80)
        d.text((1370, yy), line, font=FONT_SMALL, fill=color)
        yy += 58
    d.text((120, 925), "注：左侧历史快照与右侧工程门禁不是同一主结果面；正式投稿需以统一 frozen bundle 生成主结果。", font=FONT_SMALL, fill=(75, 85, 95))
    out = FIG_DIR / "fig3_metrics_gate_overview.png"
    img.save(out, dpi=(220, 220))
    return out


def make_fig4():
    img, d = base_canvas("图4  典型案例与失败模式")
    headers = ["案例", "题型", "观察", "论文解释边界"]
    cols = [80, 300, 520, 1040, 1700]
    y0 = 180
    row_h = 115
    d.rectangle((cols[0], y0, cols[-1], y0 + row_h), fill=(226, 235, 244), outline=(150, 165, 180), width=2)
    for i, h in enumerate(headers):
        draw_centered_text(d, (cols[i], y0, cols[i + 1], y0 + row_h), h, FONT_H, max_chars=9)
    rows = [
        ("Q02", "事实题", "1 次调用，约 200 ms，命中 n0002，回答 002196", "低成本事实定位案例"),
        ("Q33", "拒答题", "Transformer encoder 信息不在图谱中，系统拒答", "支持小样本拒答行为信号"),
        ("Q49", "假前提", "纠正“储能降低产量”前提，9023 → 10903 吨", "支持纠错行为信号"),
        ("Q10/Q11", "图表描述", "检索到部分相关节点，但关键词覆盖不足", "词面失配与图表表达风险"),
        ("Q15/Q16", "列表/指标", "漏掉 figure 或 metric 目标节点", "列表边界与指标节点盲区"),
        ("Q31", "复杂比较", "命中证据但 100 次调用、约 1.6M tokens、137 s", "答对但成本异常，提示搜索发散"),
    ]
    y = y0 + row_h
    for r_idx, row in enumerate(rows):
        fill = (250, 252, 253) if r_idx % 2 == 0 else (243, 247, 250)
        d.rectangle((cols[0], y, cols[-1], y + row_h), fill=fill, outline=(185, 195, 205), width=1)
        for i, cell in enumerate(row):
            maxc = [8, 8, 22, 22][i]
            draw_centered_text(d, (cols[i] + 8, y + 8, cols[i + 1] - 8, y + row_h - 8), cell, FONT_SMALL, max_chars=maxc)
        y += row_h
    out = FIG_DIR / "fig4_cases_failures.png"
    img.save(out, dpi=(220, 220))
    return out


def set_east_asia(run, font_name="SimSun"):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def set_cell_text(cell, text, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    r = p.add_run(str(text))
    set_east_asia(r, "SimSun")
    r.font.size = Pt(9)
    r.bold = bold
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table):
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for name in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        border = OxmlElement(f"w:{name}")
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "4")
        border.set(qn("w:space"), "0")
        border.set(qn("w:color"), "B8C2CC")
        borders.append(border)
    tblPr.append(borders)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def add_para(doc, text="", style=None, first_line=True):
    p = doc.add_paragraph(style=style)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_east_asia(r, "SimSun")
    r.font.size = Pt(10.5)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_heading(level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_east_asia(run, "SimHei")
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.bold = True
    if level == 1:
        run.font.size = Pt(14)
    else:
        run.font.size = Pt(12)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_east_asia(r, "SimSun")
    r.font.size = Pt(9)
    return p


def add_picture(doc, path, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(5.9))
    add_caption(doc, caption)


def add_table(doc, headers, rows, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(caption)
    set_east_asia(r, "SimSun")
    r.font.size = Pt(9)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    set_table_borders(table)
    for i, h in enumerate(headers):
        shade_cell(table.rows[0].cells[i], "DDEBF7")
        set_cell_text(table.rows[0].cells[i], h, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_text(cells[i], val, align=WD_ALIGN_PARAGRAPH.LEFT if i > 0 else WD_ALIGN_PARAGRAPH.CENTER)
    doc.add_paragraph()


def build_docx(figs):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.54)
    sec.bottom_margin = Cm(2.54)
    sec.left_margin = Cm(2.8)
    sec.right_margin = Cm(2.6)

    styles = doc.styles
    styles["Normal"].font.name = "SimSun"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "SimSun")
    styles["Normal"].font.size = Pt(10.5)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = title.add_run("面向中文复杂长文档问答的可追溯 GraphRAG 系统设计与评测")
    set_east_asia(tr, "SimHei")
    tr.bold = True
    tr.font.size = Pt(18)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("作者：[待补]    单位：[待补]")
    set_east_asia(r, "SimSun")
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.2
    r = p.add_run("摘要：")
    set_east_asia(r, "SimHei")
    r.bold = True
    r.font.size = Pt(10.5)
    abstract = (
        "中文复杂长文档问答不仅要求系统从长篇文档中找到相关内容，还要求回答能够回溯到明确证据，并在证据不足或问题前提错误时保持拒答和纠错能力。"
        "围绕这一任务，本文构建并评估一条由知识图谱证据组织、多路径检索、预算约束控制和引用校验生成组成的可追溯 GraphRAG 系统链路。"
        "为避免不同证据面混写，本文将生成前最终证据集合定义为检索主结果面，将早期路径命中、重排前候选集合、回答支持指标和工程门禁结果分别作为诊断面、支撑面和工程面解释。"
        "基于现有历史评测材料，系统在一篇 31 页中文文档、439 个节点、1901 条边和 50 个问题构成的评测集中，显示出高召回、引用可追溯、拒答和假前提处理等正面信号；同时也暴露出低精度、图表与列表类问题盲区，以及复杂问题上的长尾成本。"
        "当前工程门禁结果为 NO-GO，说明系统在部分 final-context 指标和尾延迟上已有通过项，但关键词覆盖和后端验证仍未达门限。"
        "结果表明，知识图谱增强检索与引用校验生成有助于组织中文长文档问答中的证据链，但该路线的应用价值依赖于明确的检索精度边界、成本边界和证据治理边界。"
    )
    r = p.add_run(abstract)
    set_east_asia(r, "SimSun")
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    r = p.add_run("关键词：")
    set_east_asia(r, "SimHei")
    r.bold = True
    r.font.size = Pt(10.5)
    r = p.add_run("中文长文档问答；GraphRAG；知识图谱增强检索；可追溯生成；拒答；证据治理")
    set_east_asia(r, "SimSun")
    r.font.size = Pt(10.5)

    p = doc.add_paragraph()
    r = p.add_run("中图分类号：[待补]    文献标识码：A")
    set_east_asia(r, "SimSun")
    r.font.size = Pt(10)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Traceable GraphRAG System Design and Evaluation for Chinese Complex Long-document Question Answering")
    r.font.name = "Times New Roman"
    r.bold = True
    r.font.size = Pt(13)
    p = doc.add_paragraph()
    r = p.add_run("Abstract: ")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    en_abs = (
        "Chinese complex long-document question answering requires not only retrieving relevant evidence, but also producing answers that remain traceable, refuse unsupported questions, and correct false premises when document evidence allows. "
        "This paper reports a traceable GraphRAG system organized around knowledge-graph evidence, multi-path retrieval, budget-constrained control, and citation-verified generation. "
        "The manuscript separates the final context used for generation from diagnostic retrieval surfaces and engineering release gates. "
        "Historical evidence from a 31-page Chinese document with 439 nodes, 1901 edges, and 50 questions indicates positive signals in recall, citation traceability, refusal, and false-premise handling, while also revealing low precision, figure/list blind spots, and long-tail cost. "
        "The current engineering release gate remains NO-GO, so the results should be interpreted as application-oriented system evidence rather than a finalized academic headline result."
    )
    r = p.add_run(en_abs)
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    p = doc.add_paragraph()
    r = p.add_run("Key words: ")
    r.bold = True
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)
    r = p.add_run("Chinese long-document QA; GraphRAG; knowledge-graph-enhanced retrieval; traceable generation; refusal; evidence governance")
    r.font.name = "Times New Roman"
    r.font.size = Pt(10)

    add_heading(doc, "1 引言", 1)
    for para in [
        "中文长文档问答广泛存在于论文、行业报告、技术说明书和项目材料等应用场景中。与短文本事实问答不同，长文档中的用户问题往往涉及章节、图表、指标、列表和跨段落关系。系统如果只能生成一段流畅回答，却无法说明回答来自哪些证据、证据是否确实被检索到、问题证据不足时是否应拒答，就很难满足复杂文档使用中的可信需求。",
        "复杂长文档问答的关键瓶颈在于证据组织。文档中的有效信息可能分散在标题、正文、图题、列表项、指标节点和相邻关系中，问题表述也常与原文表述不完全一致。单一路径检索容易遗漏关键节点；无限制地扩展检索范围，又会带入大量无关内容，增加生成阶段的上下文负担和误答风险。",
        "检索增强生成将外部检索内容引入生成模型，为知识密集型任务提供了重要思路[1]。进一步地，图结构可用于表达实体、关系和局部结构，近期 GraphRAG 相关研究也将图结构引入检索增强生成流程[2-4]。本文不将工作定位为算法竞赛式方法，而关注一条面向中文复杂长文档问答的可审计系统链路如何被组织、评测和解释。",
        "本文主要贡献体现在三方面。第一，围绕中文复杂长文档问答，给出一条由知识图谱证据、多路径检索、预算约束控制和引用校验生成组成的系统链路。第二，将 final-context 定义为检索主结果面，并区分 path_hits、pre-rerank、answer-support 和 release gate 等不同证据面。第三，基于现有历史评测和当前工程门禁材料，分析系统在引用可追溯、拒答、假前提纠错和质量-成本权衡上的行为特征，同时保留低精度、图表/列表盲区、长尾成本和外部有效性不足等限制。",
    ]:
        add_para(doc, para)

    add_heading(doc, "2 相关研究", 1)
    for para in [
        "检索增强生成通过在生成前引入外部知识，缓解大模型仅依赖参数记忆带来的知识更新和证据追溯问题。Lewis 等[1]提出的 RAG 框架将参数化生成模型与非参数化检索记忆结合，强调可访问外部知识对知识密集型任务的价值。对长文档问答而言，检索增强生成的关键不只是召回相关片段，还包括候选证据筛选、上下文预算控制和回答证据一致性。",
        "GraphRAG 相关研究进一步强调图结构在检索增强生成中的作用。已有综述将 GraphRAG 视为由查询处理、检索、组织、生成和数据源等部分构成的系统流程[2]，也有研究从图索引、图引导检索和图增强生成角度总结该方向[3]。LEGO-GraphRAG 则从模块化视角讨论 GraphRAG 工作流、设计空间和质量-成本权衡[4]。这些研究为本文的系统拆分提供了方法背景。",
        "本文与上述研究的差异在于，本文不声称提出新的知识图谱构建算法，而是将知识图谱证据、多路径检索、final-context 主结果面、引用校验生成和拒答/纠错行为放入同一证据治理框架。该定位更接近中文复杂长文档问答的系统设计与评测，而非通用算法竞争。",
        "[需要补文献：正式投稿版仍需补充中文核心或同类期刊中关于中文长文档问答、知识图谱问答系统、可追溯生成和 RAG 评测的经核验文献。]",
    ]:
        add_para(doc, para)

    add_heading(doc, "3 系统设计", 1)
    add_heading(doc, "3.1 任务定义", 2)
    add_para(doc, "本文讨论的任务是在给定中文长文档及其结构化证据的条件下，对用户问题进行文档证据约束下的回答。系统需要完成三类行为：当文档中存在支持证据时给出回答；当文档证据不足时拒答；当问题包含错误前提且文档存在反证时纠正前提。该任务不追求开放域外推，重点在回答与文档证据之间的可追溯关系。")
    add_heading(doc, "3.2 总体框架", 2)
    add_para(doc, "系统链路由文档知识图谱证据组织、查询预处理、多路径检索、重排与 final-context 选择、冗余/来源预算过滤、预算约束控制、引用校验生成以及拒答/假前提纠错组成。问题进入系统后，首先被转换为适合检索的查询表达；随后系统从知识图谱节点、关系和候选文本中获取证据；候选证据经过重排和过滤形成生成前最终证据集合；生成阶段只基于该证据集合组织回答、引用、拒答或纠错。")
    add_picture(doc, figs[0], "图1  面向中文长文档问答的可追溯 GraphRAG 系统架构")
    add_heading(doc, "3.3 知识图谱证据组织", 2)
    add_para(doc, "知识图谱证据组织的作用是将长文档中的章节、图表、指标和相邻关系转化为可检索、可引用的节点与边。对于图表解释、列表枚举和成本比较等问题，单纯段落相似度检索可能无法稳定定位目标证据。知识图谱证据使系统能够围绕节点类型、局部邻居和跨节点关系组织候选证据。")
    add_para(doc, "这一模块的边界也需要明确。知识图谱只是在文档范围内重组证据，并不生成文档外知识。若图表节点、指标节点或列表边界在图谱构建阶段不完整，后续检索仍可能遗漏关键证据。历史失败样本 Q15 和 Q16 已提示，图表和指标类节点仍是系统薄弱环节。")
    add_heading(doc, "3.4 多路径检索与证据面治理", 2)
    add_para(doc, "多路径检索用于降低单一路径漏检风险。系统通过不同检索入口形成候选证据，再进入重排、过滤和 final-context 选择阶段。本文将 path_hits 定义为早期路径命中，将 pre-rerank 定义为重排前候选集合，将 final-context 定义为真正送入生成阶段的最终证据集合。只有 final-context 可作为论文检索主结果面，path_hits 和 pre-rerank 仅用于诊断证据在哪个阶段获得或丢失。")
    add_picture(doc, figs[1], "图2  证据面分层与论文写作准入规则")
    add_heading(doc, "3.5 预算约束控制与引用校验生成", 2)
    add_para(doc, "预算约束控制用于限制工具调用、候选扩张、上下文长度、延迟和 token 成本。历史样本 Q31 表明，系统可以命中目标证据，但代价可能达到 100 次工具调用、约 1.6M tokens 和 137 秒延迟。此类长尾案例说明，复杂问题上的搜索发散是系统设计必须处理的风险。")
    add_para(doc, "引用校验生成要求回答中的节点引用来自已检索证据集合。历史快照中，回答引用节点真实性和引用来自已检索证据均为 1.000，支持系统在引用可追溯方面具有正面信号。需要强调的是，引用可追溯不等同于语义完全正确，也不等同于关键词完全覆盖。")

    add_heading(doc, "4 实验设计与证据来源", 1)
    add_heading(doc, "4.1 数据与题型", 2)
    add_para(doc, "现有历史综合评测基于一篇 31 页中文文档开展，该文档被组织为 439 个节点和 1901 条边的知识图谱。评测集包含 50 个问题，覆盖事实题、聚合题、数值题、多跳题、拒答题、噪声题和假前提题七类任务。该设置不是大规模通用基准，而是围绕中文复杂长文档问答中容易失效的行为设计。")
    add_table(doc, ["对象", "内容"], [
        ("文档规模", "31 页中文文档"),
        ("知识图谱规模", "439 个节点，1901 条边"),
        ("问题规模", "50 个问题"),
        ("题型", "事实、聚合、数值、多跳、拒答、噪声、假前提"),
        ("样本边界", "T5/T6/T7 样本量分别为 10、5、3，需谨慎解释"),
    ], "表1  历史评测对象与题型")
    add_heading(doc, "4.2 指标与证据面", 2)
    add_para(doc, "检索侧指标包括宏平均召回、宏平均精度和 F1 等，用于判断系统是否找到目标证据以及候选证据是否过宽。回答支持指标包括引用节点真实性、引用来自已检索证据、幻觉率和关键词覆盖等，用于判断回答是否建立在可检查证据之上。性能与成本指标包括工具调用、token、延迟和成本，用于分析质量-成本权衡。")
    add_para(doc, "按照指标契约，final-context 是论文默认检索主结果面。path_hits 和 pre-rerank 只能作为诊断面使用，不能进入摘要、引言和主结果表。历史评测快照属于 historical evidence；当前 release gate 属于 engineering gate；二者均不能替代未来统一冻结后的 current frozen evidence。")

    add_heading(doc, "5 实验结果与分析", 1)
    add_heading(doc, "5.1 检索与回答支持", 2)
    add_para(doc, "历史快照显示，系统呈现高召回、低精度的检索倾向。REPORT_FINAL.md 记录的 Recall_macro 为 0.82，Precision_macro 为 0.078，平均检索集合规模约为 40，平均真值节点规模约为 2。这说明系统在历史评测面上更偏向避免漏掉证据，同时带入较多非目标节点。")
    add_para(doc, "回答支持面给出了较强的可追溯信号。历史快照记录 node_id_validity=1.000、cited_in_retrieved=1.000、hallucination_rate=0，说明回答中出现的节点引用均真实存在，并且来自系统实际检索过的证据集合。与此同时，keyword_coverage=0.799，且存在一次 forbidden violation，说明引用可追溯并不自动等于答案在词面上完全贴合真值。")
    add_picture(doc, figs[2], "图3  历史快照与当前工程门禁的主要信号")
    add_table(doc, ["指标/结果", "数值", "证据类别", "解释边界"], [
        ("Recall_macro", "0.82", "historical evidence", "历史快照高召回信号，非最终主结果"),
        ("Precision_macro", "0.078", "historical evidence", "提示候选过宽和低精度"),
        ("node_id_validity", "1.000", "answer-support", "回答引用节点真实存在"),
        ("cited_in_retrieved", "1.000", "answer-support", "回答引用来自已检索证据"),
        ("keyword_coverage", "0.799", "answer-support", "词面代理指标，不等同语义正确率"),
        ("release gate verdict", "NO-GO", "engineering gate", "工程状态，不能替代学术主结果"),
    ], "表2  当前可用指标与解释边界")
    add_heading(doc, "5.2 拒答、噪声鲁棒与假前提处理", 2)
    add_para(doc, "历史快照中的 T5、T6、T7 为系统行为可靠性提供了小样本观察。T5 拒答题为 10/10，T6 噪声题为 5/5，T7 假前提题为 3/3。该结果说明，在这组历史样本中，系统没有一味生成答案，而是在文档外问题、噪声问题和错误前提问题上显示出保守行为。")
    add_para(doc, "这部分结论必须与样本量一起解释。T5、T6、T7 的样本量分别只有 10、5 和 3，且假前提题内部还包括拒答式假前提和纠正式假前提。正式投稿版需要补 Wilson 或 exact binomial 置信区间，并将拒答、噪声鲁棒和假前提纠错分别解释。")
    add_heading(doc, "5.3 质量-成本权衡", 2)
    add_para(doc, "历史快照同时暴露出质量和成本之间的张力。多跳问题上，系统通过更大范围探索获得证据，但工具调用、token 和延迟同步增加。Q31 是最典型的长尾案例：该题最终命中了所需证据，但伴随 100 次工具调用、约 1.6M tokens、137 秒延迟，并占总成本的 35.5%。报告将其解释为探索发散，而不是循环死锁。")
    add_heading(doc, "5.4 案例与失败分析", 2)
    add_para(doc, "Q02 可作为事实回答案例。问题是“论文编号是什么？”，系统以 1 次工具调用、约 200 ms 命中节点 n0002，并回答“002196”。Q33 可作为拒答案例，系统确认知识图谱无相关实体后明确拒答。Q49 可作为假前提纠错案例，系统依据文档证据反驳“储能反而降低产量”的错误前提，指出实际从 9023 吨提升至 10903 吨。")
    add_para(doc, "失败案例更能说明系统边界。Q10 和 Q11 属于图表描述问题，主要体现为词面失配；Q15 错过 3 个 figure 节点，属于列表聚合真错；Q16 错过目标 metric 节点，反映列表边界和指标节点处理不足；Q31 则显示复杂比较问题可能出现搜索发散。")
    add_picture(doc, figs[3], "图4  典型案例与失败模式")
    add_heading(doc, "5.5 待补实验", 2)
    for gap in [
        "[需要补证据：统一 frozen rerun 后的 final-context 主结果表。]",
        "[需要补证据：预算匹配 baseline。]",
        "[需要补证据：关键模块消融。]",
        "[需要补证据：T5/T6/T7 的置信区间。]",
        "[需要补证据：跨文档迁移、人评、统计分析和 LLM 稳定性协议。]",
    ]:
        add_para(doc, gap)

    add_heading(doc, "6 讨论", 1)
    for title2, paras in [
        ("6.1 系统证明了什么", [
            "在当前证据范围内，系统证明的是一条可追溯问答链路可以被组织起来，并且能够用多层指标检查其行为。历史快照显示，系统在引用节点真实性、引用来自已检索证据、拒答、噪声问题和假前提处理上有正面信号。案例层面，Q02、Q33 和 Q49 展示了事实回答、拒答和假前提纠错的可检查过程。",
            "这一证明的重点不是性能结论已经充分建立，而是回答行为可以被拆开检查。对中文复杂长文档问答而言，用户不仅关心答案，还关心答案是否有证据、是否越界、是否在错误问题上保持保守。",
        ]),
        ("6.2 系统没有证明什么", [
            "当前材料没有证明系统相对既有路线已经建立比较优势。预算匹配 baseline 尚未完成，不能判断完整 GraphRAG 链路相对 BM25、向量检索、混合检索或固定流程 RAG 的净收益。关键模块消融尚未完成，也不能判断知识图谱证据、多路径检索、重排、预算控制或引用校验分别贡献了多少。",
            "当前材料也没有证明外部有效性。历史评测主要来自单篇中文文档，50 题规模能够观察行为，但不能外推到中文长文档问答总体。T5、T6、T7 样本量较小，拒答与纠错结论需要配置信区间或扩样。当前 release gate 为 NO-GO，说明工程状态也不能作为系统已经准备好的证据。",
        ]),
        ("6.3 应用价值与边界", [
            "对中文复杂长文档问答场景而言，可追溯 GraphRAG 系统的应用意义在于提供一种可检查的回答流程。它把回答、证据、拒答和纠错放在同一链路中，使使用者能够看到系统依据什么回答、在什么情况下停下、在什么情况下纠正问题前提。",
            "这种应用意义必须和局限一起理解。系统当前仍存在低精度、图表/列表盲区、长尾成本、单文档评测、人工标注边界和工程门禁 NO-GO 等问题。正式投稿版如果不能补齐 baseline、ablation、CI、迁移和人评，结论应停留在具有应用潜力和进一步验证价值，不能升级为有效性已经充分建立。",
        ]),
    ]:
        add_heading(doc, title2, 2)
        for para in paras:
            add_para(doc, para)

    add_heading(doc, "7 结论", 1)
    add_para(doc, "本文围绕中文复杂长文档问答，整理并评估了一条由知识图谱证据、多路径检索、预算约束控制和引用校验生成构成的可追溯 GraphRAG 系统链路。现有材料显示，该链路能够把证据获取、回答生成、拒答和假前提纠错纳入同一可检查流程；历史快照也提示系统在引用可追溯和行为保守性方面具有正面信号。")
    add_para(doc, "这一路线仍没有完成投稿级证据闭环。当前缺少统一 frozen final-context 主结果、预算匹配 baseline、关键模块消融、鲁棒性置信区间、跨文档迁移、人评和外部文献台账；当前工程门禁也处于 NO-GO 状态。因此，更稳妥的结论是：可追溯 GraphRAG 为中文复杂长文档问答提供了一条值得继续验证的系统化方案，其价值成立依赖精度边界、成本边界和证据治理边界共同收敛。")

    add_heading(doc, "参考文献", 1)
    refs = [
        "Lewis P, Perez E, Piktus A, et al. Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks[EB/OL]. arXiv:2005.11401, 2020.",
        "Han H, Wang Y, Shomer H, et al. Retrieval-Augmented Generation with Graphs (GraphRAG)[EB/OL]. arXiv:2501.00309, 2024.",
        "Peng B, Zhu Y, Liu Y, et al. Graph Retrieval-Augmented Generation: A Survey[EB/OL]. arXiv:2408.08921, 2024.",
        "Cao Y, Gao Z, Li Z, et al. LEGO-GraphRAG: Modularizing Graph-based Retrieval-Augmented Generation for Design Space Exploration[EB/OL]. arXiv:2411.05844, 2024.",
        "[需要补文献] 中文核心或同类期刊中关于中文长文档问答、知识图谱问答系统、可追溯生成和 RAG 评测的经核验文献。",
    ]
    for i, ref in enumerate(refs, 1):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.74)
        r = p.add_run(f"[{i}] {ref}")
        set_east_asia(r, "SimSun")
        r.font.size = Pt(9)

    doc.save(DOCX_PATH)


def main():
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    figs = [make_fig1(), make_fig2(), make_fig3(), make_fig4()]
    build_docx(figs)
    print(DOCX_PATH)
    for f in figs:
        print(f)


if __name__ == "__main__":
    main()
